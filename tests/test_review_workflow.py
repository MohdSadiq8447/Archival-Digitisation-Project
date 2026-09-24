from __future__ import annotations

import asyncio
import json
from dataclasses import asdict
from pathlib import Path
from types import SimpleNamespace
from typing import Any, cast

import pandas as pd
import pytest
import yaml
from openpyxl import load_workbook
from PIL import Image

from census_extractor.metadata import MetadataRegistry
from census_extractor.pipeline.runner import PipelineRunner
from census_extractor.review import (
    REVIEW_HEADERS,
    ReviewCell,
    ReviewPackageBuilder,
    ReviewWorkbookCompiler,
    sha256_file,
)
from census_extractor.schemas import SchemaRegistry, TableSchema
from census_extractor.workflow import RemainingUPWorkflow


def _settings_path(project_config) -> Path:
    return project_config.data_dir / "workflow_remaining_up.yaml"


def test_remaining_workflow_selects_exact_scope_and_deterministic_order(project_config):
    workflow = RemainingUPWorkflow(project_config, _settings_path(project_config))
    documents = workflow.documents
    assert len(documents) == 142
    assert len({document.district for document in documents}) == 48
    assert {
        format_id: sum(document.format_id == format_id for document in documents)
        for format_id in ("format_001", "format_002", "format_003")
    } == {"format_001": 48, "format_002": 48, "format_003": 46}
    assert not {
        "Aligarh",
        "Allahabad",
        "Almora",
        "Azamgarh",
        "Bahraich",
    }.intersection(document.district for document in documents)
    by_district: dict[str, list[str]] = {}
    for document in documents:
        by_district.setdefault(document.district, []).append(document.format_id)
    assert by_district["Mirzapur"] == ["format_001", "format_002"]
    assert by_district["Varanasi"] == ["format_001", "format_002"]
    assert documents == sorted(
        documents,
        key=lambda item: (
            item.district.casefold(),
            ["format_001", "format_002", "format_003"].index(item.format_id),
        ),
    )


def test_v3_workflow_reaches_provisional_complete_without_review_wait(
    project_config, tmp_path, monkeypatch
):
    config = project_config.with_overrides(output_dir=tmp_path)
    workflow = RemainingUPWorkflow(config, _settings_path(project_config))
    state: dict[str, Any] = {"stage": "INITIAL"}
    ledger_path = workflow.automatic_ledger
    ledger_path.parent.mkdir(parents=True)
    ledger_path.write_text("version: 2\n", encoding="utf-8")
    extraction_manifest = (
        config.runs_dir / workflow.settings.extraction_run_id / "manifest.json"
    )
    extraction_manifest.parent.mkdir(parents=True)
    extraction_manifest.write_text('{"results": {}}', encoding="utf-8")

    monkeypatch.setattr(workflow, "_load_or_create_state", lambda: state)
    monkeypatch.setattr(workflow, "_run_geometry_stage", lambda _state: [])
    monkeypatch.setattr(
        workflow, "_run_extraction_stage", lambda _state, _documents: []
    )
    saved_stages: list[str] = []
    monkeypatch.setattr(
        workflow,
        "_save_state",
        lambda value: saved_stages.append(str(value["stage"])),
    )

    class FakeCorrector:
        def __init__(self, *_args, **_kwargs):
            pass

        async def run_async(self, **_kwargs):
            return SimpleNamespace(
                ledger_path=ledger_path,
                included_pdf_ids=[item.pdf_id for item in workflow.documents],
                exclusions=[],
                corrected_cells=2,
                confirmed_cells=3,
                unresolved_cells=1,
                cache_hits=4,
                cache_misses=5,
            )

    class FakePostprocessor:
        def __init__(self, *_args, **_kwargs):
            pass

        def run_automatic(self, **_kwargs):
            workflow.final_root.mkdir(parents=True)
            return SimpleNamespace(
                output_root=workflow.final_root,
                csv_files=[Path("one.csv")] * len(workflow.documents),
                reviewed_unique_cells=6,
                corrected_cells=2,
                unresolved_cells=1,
            )

    class FakePackageBuilder:
        def __init__(self, *_args, **_kwargs):
            pass

        def build(self, **_kwargs):
            return {"selected_pdfs": len(workflow.documents), "unique_source_cells": 9}

    monkeypatch.setattr(
        "census_extractor.workflow.AutomaticTranscriptionCorrector", FakeCorrector
    )
    monkeypatch.setattr("census_extractor.workflow.CSVPostprocessor", FakePostprocessor)
    monkeypatch.setattr(
        "census_extractor.workflow.CodexVerificationPackageBuilder", FakePackageBuilder
    )
    monkeypatch.setattr(
        workflow,
        "_merge_outputs",
        lambda _documents: {"format_001": {"sources": 48, "rows": 1}},
    )
    report = workflow.workflow_root / "WORKFLOW_QA_REPORT.md"

    def write_report(*_args, **_kwargs):
        report.parent.mkdir(parents=True, exist_ok=True)
        report.write_text("provisional", encoding="utf-8")
        return report

    monkeypatch.setattr(workflow, "_write_workflow_report", write_report)
    result = workflow.run()
    assert result.status == "PROVISIONAL_COMPLETE"
    assert state["stage"] == "PROVISIONAL_COMPLETE"
    assert "WAITING_FOR_REVIEW" not in saved_stages
    assert not (workflow.workflow_root / "SOURCE_REVIEW.xlsx").exists()
    assert state["automatic_correction"]["unresolved_cells"] == 1
    assert state["codex_verification_package"]["selected_pdfs"] == 142


def test_v3_workflow_noop_and_structural_blocking_states(
    project_config, tmp_path, monkeypatch
):
    workflow = RemainingUPWorkflow(
        project_config.with_overrides(output_dir=tmp_path),
        _settings_path(project_config),
    )
    monkeypatch.setattr(workflow, "_save_state", lambda _state: None)

    complete: dict[str, Any] = {"stage": "PROVISIONAL_COMPLETE"}
    monkeypatch.setattr(workflow, "_load_or_create_state", lambda: complete)
    result = workflow.run()
    assert result.status == "PROVISIONAL_COMPLETE"
    assert "no OCR" in result.message

    geometry_state: dict[str, Any] = {"stage": "INITIAL"}
    monkeypatch.setattr(workflow, "_load_or_create_state", lambda: geometry_state)
    monkeypatch.setattr(
        workflow,
        "_run_geometry_stage",
        lambda _state: [{"pdf_id": "fixture", "reason": "missing panel"}],
    )
    result = workflow.run()
    assert result.status == "BLOCKED_GEOMETRY"
    assert geometry_state["blocking_failures"][0]["pdf_id"] == "fixture"

    extraction_state: dict[str, Any] = {"stage": "GEOMETRY_COMPLETE"}
    monkeypatch.setattr(workflow, "_load_or_create_state", lambda: extraction_state)
    monkeypatch.setattr(workflow, "_run_geometry_stage", lambda _state: [])
    monkeypatch.setattr(
        workflow,
        "_run_extraction_stage",
        lambda _state, _documents: [
            {"pdf_id": "fixture", "reason": "operational error"}
        ],
    )
    result = workflow.run()
    assert result.status == "BLOCKED_EXTRACTION"
    assert extraction_state["stage"] == "BLOCKED_EXTRACTION"


def _review_fixture(tmp_path: Path, *, status: str = "PENDING") -> tuple[Path, Path, Path]:
    manifest = tmp_path / "runs" / "source" / "manifest.json"
    manifest.parent.mkdir(parents=True)
    manifest.write_text('{"results": {}}', encoding="utf-8")
    crop = tmp_path / "review" / "crops" / "fixture" / "r0" / "01_sl_no.png"
    crop.parent.mkdir(parents=True)
    crop.write_bytes(b"png fixture")
    workbook_path = tmp_path / "review" / "SOURCE_REVIEW.xlsx"
    index_path = tmp_path / "review" / "review_index.json"
    cell = ReviewCell(
        review_id="cell-1",
        pdf_id="fixture",
        district="Fixture",
        format_id="format_001",
        scope="row",
        row_index=0,
        parent_row_index=None,
        subrow_index=None,
        variable="sl_no",
        extracted_value="=literal printed text",
        review_status=status,
        corrected_value="",
        reason="",
        source_page=1,
        panel_id="civic_anchor",
        bbox=(1, 2, 30, 20),
        crop_path="crops/fixture/r0/01_sl_no.png",
        crop_sha256=sha256_file(crop),
    )
    index = {
        "version": 1,
        "source_run_id": "source",
        "source_manifest": str(manifest),
        "source_manifest_sha256": sha256_file(manifest),
        "review_workbook": workbook_path.name,
        "table_summary": [
            {
                "pdf_id": "fixture",
                "district": "Fixture",
                "format_id": "format_001",
                "rows": 1,
                "parents": 1,
                "unique_cells": 1,
            }
        ],
        "cells": [
            {
                **asdict(cell),
                "bbox": list(cell.bbox),
            }
        ],
    }
    index_path.write_text(json.dumps(index, indent=2), encoding="utf-8")
    ReviewPackageBuilder._write_workbook(
        workbook_path,
        [cell],
        index["table_summary"],
        manifest_sha=index["source_manifest_sha256"],
        index_sha=sha256_file(index_path),
    )
    return workbook_path, index_path, crop


def test_review_workbook_is_formula_safe_styled_and_pauses_while_pending(tmp_path):
    workbook_path, index_path, _ = _review_fixture(tmp_path)
    workbook = load_workbook(workbook_path, data_only=False)
    try:
        assert {"Instructions", "Summary", "Cells", "_Metadata"}.issubset(
            workbook.sheetnames
        )
        sheet = workbook["Cells"]
        assert [cell.value for cell in sheet[1]] == REVIEW_HEADERS
        extracted_column = REVIEW_HEADERS.index("extracted_value") + 1
        assert sheet.cell(2, extracted_column).value == "=literal printed text"
        assert sheet.cell(2, extracted_column).data_type == "s"
        assert sheet.freeze_panes == "A2"
        assert sheet.tables["SourceReviewCells"].ref == "A1:R2"
        assert sheet.cell(2, REVIEW_HEADERS.index("crop_path") + 1).hyperlink is not None
    finally:
        workbook.close()
    result = ReviewWorkbookCompiler().compile(
        workbook_path=workbook_path,
        index_path=index_path,
        ledger_path=tmp_path / "ledger.yaml",
    )
    assert result is None
    assert not (tmp_path / "ledger.yaml").exists()


def test_review_workbook_compiles_correction_and_guards_immutable_fields(tmp_path):
    workbook_path, index_path, _ = _review_fixture(tmp_path)
    workbook = load_workbook(workbook_path)
    sheet = workbook["Cells"]
    sheet.cell(2, REVIEW_HEADERS.index("review_status") + 1, "CORRECTED")
    sheet.cell(2, REVIEW_HEADERS.index("corrected_value") + 1, "1")
    sheet.cell(2, REVIEW_HEADERS.index("reason") + 1, "verified against source crop")
    workbook.save(workbook_path)
    workbook.close()
    ledger_path = tmp_path / "ledger.yaml"
    result = ReviewWorkbookCompiler().compile(
        workbook_path=workbook_path,
        index_path=index_path,
        ledger_path=ledger_path,
    )
    assert result is not None
    assert result.reviewed_cells == 1
    assert result.corrected_cells == 1
    ledger = yaml.safe_load(ledger_path.read_text(encoding="utf-8"))
    assert ledger["version"] == 2
    assert ledger["tables"]["fixture"]["reviews"][0]["verified_value"] == "1"

    workbook = load_workbook(workbook_path)
    workbook["Cells"].cell(2, REVIEW_HEADERS.index("variable") + 1, "town_name")
    workbook.save(workbook_path)
    workbook.close()
    with pytest.raises(ValueError, match="immutable field"):
        ReviewWorkbookCompiler().compile(
            workbook_path=workbook_path,
            index_path=index_path,
            ledger_path=tmp_path / "second.yaml",
        )
    assert not (tmp_path / "second.yaml").exists()


def test_review_geometry_uses_strict_column_and_child_bands():
    panel = {
        "body_bbox": [0, 100, 300, 400],
        "column_spans": {
            "3": {"x_start": 80, "x_end": 160},
            "4": {"x_start": 160, "x_end": 220},
        },
        "rows": [{"row_index": 0, "bbox": [0, 120, 300, 220]}],
        "hierarchy": {
            "parents": [
                {
                    "parent_row_index": 0,
                    "subrows": [
                        {"subrow_index": 0, "bbox": [80, 130, 220, 160]},
                        {"subrow_index": 1, "bbox": [80, 165, 220, 195]},
                    ],
                }
            ]
        },
    }
    assert ReviewPackageBuilder._column_span(panel, 3) == (80, 160)
    schema = SimpleNamespace(
        hierarchy=SimpleNamespace(child_variables=["hospitals_dispensaries", "med_beds"])
    )
    assert ReviewPackageBuilder._row_y_span(
        panel,
        cast(TableSchema, schema),
        "med_beds",
        None,
        0,
        1,
    ) == (165, 195)


def test_review_builder_creates_one_strict_crop_per_flat_source_cell(
    project_config, tmp_path
):
    schema = SchemaRegistry(project_config.schemas_dir).require("format_001")
    document = MetadataRegistry(project_config.metadata_path).get("agra_civic_1971")
    assert document is not None
    frame = pd.DataFrame(
        [
            {
                **{variable: "..." for variable in schema.get_all_variables()},
                "row_index": "0",
            }
        ]
    )
    panels = []
    for panel_definition in schema.panels:
        spans = {
            str(number): {"x_start": index * 20, "x_end": (index + 1) * 20}
            for index, number in enumerate(panel_definition.printed_columns)
        }
        panels.append(
            {
                "panel_id": panel_definition.panel_id,
                "page": panel_definition.page,
                "body_bbox": [0, 10, 200, 80],
                "column_spans": spans,
                "rows": [{"row_index": 0, "bbox": [0, 20, 200, 50]}],
            }
        )
    pages = [
        SimpleNamespace(width=200, height=100, image=Image.new("RGB", (200, 100), "white")),
        SimpleNamespace(width=200, height=100, image=Image.new("RGB", (200, 100), "white")),
    ]
    crop_root = tmp_path / "review" / "crops"
    cells = ReviewPackageBuilder(project_config)._table_cells(
        document,
        schema,
        frame,
        {"panels": panels},
        pages,
        crop_root,
    )
    assert len(cells) == 16
    assert len({cell.review_id for cell in cells}) == 16
    assert all((crop_root.parent / cell.crop_path).is_file() for cell in cells)
    assert cells[0].bbox == (0, 20, 20, 50)


def test_review_compiler_rejects_stale_manifest_and_changed_crop(tmp_path):
    workbook_path, index_path, crop = _review_fixture(tmp_path)
    workbook = load_workbook(workbook_path)
    workbook["Cells"].cell(2, REVIEW_HEADERS.index("review_status") + 1, "VERIFIED")
    workbook.save(workbook_path)
    workbook.close()
    manifest = Path(json.loads(index_path.read_text(encoding="utf-8"))["source_manifest"])
    manifest.write_text('{"results": {"changed": {}}}', encoding="utf-8")
    with pytest.raises(ValueError, match="manifest changed"):
        ReviewWorkbookCompiler().compile(
            workbook_path=workbook_path,
            index_path=index_path,
            ledger_path=tmp_path / "stale.yaml",
        )
    assert not (tmp_path / "stale.yaml").exists()

    workbook_path, index_path, crop = _review_fixture(tmp_path / "crop-case")
    workbook = load_workbook(workbook_path)
    workbook["Cells"].cell(2, REVIEW_HEADERS.index("review_status") + 1, "VERIFIED")
    workbook.save(workbook_path)
    workbook.close()
    crop.write_bytes(b"changed")
    with pytest.raises(ValueError, match="crop is missing or changed"):
        ReviewWorkbookCompiler().compile(
            workbook_path=workbook_path,
            index_path=index_path,
            ledger_path=tmp_path / "crop.yaml",
        )
    assert not (tmp_path / "crop.yaml").exists()


def test_sequential_helper_preserves_order_and_closes_client(project_config):
    workflow = RemainingUPWorkflow(project_config, _settings_path(project_config))
    documents = workflow.documents[:3]
    calls: list[str] = []

    class Client:
        closed = False

        async def aclose(self):
            self.closed = True

    class Runner:
        ocr_client = Client()

        async def process_pdf_async(self, path, **kwargs):
            calls.append(Path(path).stem)

    runner = Runner()
    asyncio.run(
        workflow._process_documents_sequentially(
            cast(PipelineRunner, runner), documents, is_dry_run=True
        )
    )
    assert calls == [document.pdf_id for document in documents]
    assert runner.ocr_client.closed


def test_structural_quarantine_blocks_but_transcription_quarantine_is_reviewable(tmp_path):
    report = tmp_path / "validation.json"
    record = {"status": "QUARANTINED", "exported_files": {"report": str(report)}}
    report.write_text(
        json.dumps(
            {
                "panels_complete": True,
                "alignment_complete": True,
                "findings": [
                    {"severity": "ERROR", "code": "ambiguous_ocr", "message": "cell"}
                ],
            }
        ),
        encoding="utf-8",
    )
    assert RemainingUPWorkflow._is_reviewable_result(record)
    report.write_text(
        json.dumps(
            {
                "panels_complete": True,
                "alignment_complete": False,
                "findings": [
                    {"severity": "ERROR", "code": "panel_alignment", "message": "rows"}
                ],
            }
        ),
        encoding="utf-8",
    )
    assert not RemainingUPWorkflow._is_reviewable_result(record)


def test_completed_workflow_never_waits_for_review_or_creates_workbook(
    project_config, tmp_path
):
    workflow = RemainingUPWorkflow(
        project_config.with_overrides(output_dir=tmp_path),
        _settings_path(project_config),
    )
    state = workflow._load_or_create_state()
    state["stage"] = "PROVISIONAL_COMPLETE"
    workflow._save_state(state)
    result = workflow.run()
    assert result.status == "PROVISIONAL_COMPLETE"
    assert result.review_workbook is None
    assert not list(tmp_path.rglob("SOURCE_REVIEW.xlsx"))
    assert "WAITING_FOR_REVIEW" not in result.state_path.read_text(encoding="utf-8")


def test_any_geometry_failure_blocks_paid_extraction_without_exclusions(
    project_config, tmp_path, monkeypatch
):
    workflow = RemainingUPWorkflow(
        project_config.with_overrides(output_dir=tmp_path),
        _settings_path(project_config),
    )
    blockers = [
        {"pdf_id": document.pdf_id, "reason": "geometry failure"}
        for document in workflow.documents
    ]
    monkeypatch.setattr(workflow, "_run_geometry_stage", lambda _state: blockers)

    extraction_called = False

    def extraction(_state, documents):
        nonlocal extraction_called
        extraction_called = True
        return []

    monkeypatch.setattr(workflow, "_run_extraction_stage", extraction)
    result = workflow.run()
    assert result.status == "BLOCKED_GEOMETRY"
    assert not extraction_called
    assert not list(tmp_path.rglob("SOURCE_REVIEW.xlsx"))


def test_completed_manifests_are_skipped_without_api_key(project_config, tmp_path):
    config = project_config.with_overrides(output_dir=tmp_path, novita_api_key="")
    workflow = RemainingUPWorkflow(config, _settings_path(project_config))
    geometry_manifest = config.runs_dir / workflow.settings.geometry_run_id / "manifest.json"
    extraction_manifest = config.runs_dir / workflow.settings.extraction_run_id / "manifest.json"
    geometry_manifest.parent.mkdir(parents=True)
    extraction_manifest.parent.mkdir(parents=True)
    geometry_manifest.write_text(
        json.dumps(
            {
                "results": {
                    document.pdf_id: {"status": "DRY_RUN"}
                    for document in workflow.documents
                }
            }
        ),
        encoding="utf-8",
    )
    extraction_manifest.write_text(
        json.dumps(
            {
                "results": {
                    document.pdf_id: {"status": "SUCCESS"}
                    for document in workflow.documents
                }
            }
        ),
        encoding="utf-8",
    )
    state: dict[str, object] = {}
    assert workflow._run_geometry_stage(state) == []
    assert workflow._run_extraction_stage(state) == []
    assert state["stage"] == "EXTRACTION_COMPLETE"


def test_merge_outputs_produces_48_48_46_files_without_pilot_districts(
    project_config, tmp_path
):
    config = project_config.with_overrides(output_dir=tmp_path)
    workflow = RemainingUPWorkflow(config, _settings_path(project_config))
    csv_root = workflow.final_root / "csv"
    csv_root.mkdir(parents=True)
    for document in workflow.documents:
        (csv_root / f"{document.pdf_id}.csv").write_text(
            f"row_index,pdf_id,district,format_id\n0,{document.pdf_id},{document.district},{document.format_id}\n",
            encoding="utf-8",
        )
    summary = workflow._merge_outputs()
    assert {key: value["source_csvs"] for key, value in summary.items()} == {
        "format_001": 48,
        "format_002": 48,
        "format_003": 46,
    }
    for item in summary.values():
        content = Path(item["path"]).read_text(encoding="utf-8")
        assert not any(
            pilot in content
            for pilot in ("Aligarh", "Allahabad", "Almora", "Azamgarh", "Bahraich")
        )


def test_partial_merge_always_writes_header_only_missing_formats(project_config, tmp_path):
    config = project_config.with_overrides(output_dir=tmp_path)
    workflow = RemainingUPWorkflow(config, _settings_path(project_config))
    civic = next(item for item in workflow.documents if item.format_id == "format_001")
    csv_root = workflow.final_root / "csv"
    csv_root.mkdir(parents=True)
    header = workflow._schema_header(workflow.schemas.require("format_001"))
    row = {field: "" for field in header}
    row.update(pdf_id=civic.pdf_id, district=civic.district, format_id=civic.format_id)
    pd.DataFrame([row], columns=header).to_csv(
        csv_root / f"{civic.pdf_id}.csv", index=False, lineterminator="\n"
    )
    summary = workflow._merge_outputs([civic])
    assert summary["format_001"]["source_csvs"] == 1
    assert summary["format_002"]["source_csvs"] == 0
    assert summary["format_003"]["source_csvs"] == 0
    for format_id in ("format_002", "format_003"):
        merged = Path(summary[format_id]["path"])
        assert len(merged.read_text(encoding="utf-8").splitlines()) == 1
