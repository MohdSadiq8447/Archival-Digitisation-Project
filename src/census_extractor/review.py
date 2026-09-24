"""Human source-review package generation and guarded ledger compilation."""

from __future__ import annotations

import hashlib
import json
import time
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, cast

import pandas as pd
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from census_extractor.config import PipelineConfig
from census_extractor.metadata import DocumentMetadata
from census_extractor.postprocessing import CorrectionLedger
from census_extractor.preprocessing.pdf_loader import PDFLoader
from census_extractor.schemas import SchemaRegistry, TableSchema

REVIEW_HEADERS = [
    "review_id",
    "pdf_id",
    "district",
    "format_id",
    "scope",
    "row_index",
    "parent_row_index",
    "subrow_index",
    "variable",
    "extracted_value",
    "review_status",
    "corrected_value",
    "reason",
    "source_page",
    "panel_id",
    "bbox",
    "crop_path",
    "crop_sha256",
]
EDITABLE_HEADERS = {"review_status", "corrected_value", "reason"}
IMMUTABLE_HEADERS = [header for header in REVIEW_HEADERS if header not in EDITABLE_HEADERS]


def sha256_file(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest().upper()


def _integer(value: Any) -> int:
    return int(str(value))


def _atomic_text(path: Path, text: str) -> None:
    temporary = path.with_name(f".{path.name}.{time.time_ns()}.tmp")
    temporary.write_text(text, encoding="utf-8", newline="\n")
    temporary.replace(path)


@dataclass(frozen=True, slots=True)
class ReviewCell:
    review_id: str
    pdf_id: str
    district: str
    format_id: str
    scope: str
    row_index: int | None
    parent_row_index: int | None
    subrow_index: int | None
    variable: str
    extracted_value: str
    review_status: str
    corrected_value: str
    reason: str
    source_page: int
    panel_id: str
    bbox: tuple[int, int, int, int]
    crop_path: str
    crop_sha256: str

    def workbook_row(self) -> list[Any]:
        payload = asdict(self)
        payload["bbox"] = ",".join(str(value) for value in self.bbox)
        return [payload[header] for header in REVIEW_HEADERS]


@dataclass(frozen=True, slots=True)
class ReviewPackageResult:
    workbook: Path
    index: Path
    crop_root: Path
    unique_cells: int
    index_sha256: str


@dataclass(frozen=True, slots=True)
class ReviewCompilationResult:
    ledger: Path
    reviewed_cells: int
    corrected_cells: int


class ReviewPackageBuilder:
    """Create strict cell crops plus a formula-safe Excel review workbook."""

    def __init__(self, config: PipelineConfig):
        self.config = config
        self.schemas = SchemaRegistry(config.schemas_dir)
        self.loader = PDFLoader(config.render_dpi, config.auto_deskew)

    def create(
        self,
        *,
        manifest_path: Path,
        ordered_documents: list[DocumentMetadata],
        target_dir: Path,
    ) -> ReviewPackageResult:
        manifest_path = Path(manifest_path).resolve()
        target_dir = Path(target_dir).resolve()
        workbook_path = target_dir / "SOURCE_REVIEW.xlsx"
        index_path = target_dir / "review_index.json"
        crop_root = target_dir / "crops"
        if workbook_path.exists() or index_path.exists() or crop_root.exists():
            raise FileExistsError(f"Review package already exists: {target_dir}")
        target_dir.mkdir(parents=True, exist_ok=True)
        crop_root.mkdir(parents=True)
        manifest_sha = sha256_file(manifest_path)
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        expected_ids = [document.pdf_id for document in ordered_documents]
        if set(manifest.get("results", {})) != set(expected_ids):
            raise ValueError("Extraction manifest does not match the selected review documents")

        cells: list[ReviewCell] = []
        table_summary: list[dict[str, Any]] = []
        try:
            for document in ordered_documents:
                record = manifest["results"][document.pdf_id]
                if record.get("status") not in {"SUCCESS", "QUARANTINED"}:
                    raise ValueError(
                        f"{document.pdf_id}: cannot review status {record.get('status')!r}"
                    )
                schema = self.schemas.require(document.format_id)
                frame = pd.read_csv(
                    Path(record["exported_files"]["csv"]),
                    dtype=str,
                    keep_default_na=False,
                )
                geometry_path = Path(record["exported_files"]["geometry"])
                geometry = json.loads(geometry_path.read_text(encoding="utf-8"))
                pdf_path = self.config.pdfs_dir / document.file_name
                expected_pdf_sha = str(geometry.get("source_pdf_sha256") or "").upper()
                if expected_pdf_sha and sha256_file(pdf_path) != expected_pdf_sha:
                    raise ValueError(f"{document.pdf_id}: source PDF hash changed")
                pages = self.loader.render_pdf(pdf_path)
                table_cells = self._table_cells(
                    document,
                    schema,
                    frame,
                    geometry,
                    pages,
                    crop_root,
                )
                expected_cells = self._expected_unique_cells(schema, frame)
                if len(table_cells) != expected_cells or len(
                    {cell.review_id for cell in table_cells}
                ) != expected_cells:
                    raise ValueError(
                        f"{document.pdf_id}: review coverage {len(table_cells)} != "
                        f"expected {expected_cells}"
                    )
                cells.extend(table_cells)
                table_summary.append(
                    {
                        "pdf_id": document.pdf_id,
                        "district": document.district,
                        "format_id": document.format_id,
                        "rows": len(frame),
                        "parents": (
                            frame["parent_row_index"].nunique()
                            if schema.hierarchy is not None
                            else len(frame)
                        ),
                        "unique_cells": len(table_cells),
                    }
                )
            index_payload = {
                "version": 1,
                "source_run_id": manifest_path.parent.name,
                "source_manifest": str(manifest_path),
                "source_manifest_sha256": manifest_sha,
                "review_workbook": workbook_path.name,
                "table_summary": table_summary,
                "cells": [asdict(cell) for cell in cells],
            }
            _atomic_text(
                index_path,
                json.dumps(index_payload, indent=2, ensure_ascii=False, sort_keys=True) + "\n",
            )
            index_sha = sha256_file(index_path)
            self._write_workbook(
                workbook_path,
                cells,
                table_summary,
                manifest_sha=manifest_sha,
                index_sha=index_sha,
            )
        except Exception:
            if workbook_path.exists():
                workbook_path.unlink()
            if index_path.exists():
                index_path.unlink()
            if crop_root.exists():
                import shutil

                shutil.rmtree(crop_root, ignore_errors=True)
            raise
        return ReviewPackageResult(
            workbook_path,
            index_path,
            crop_root,
            len(cells),
            index_sha,
        )

    def _table_cells(
        self,
        document: DocumentMetadata,
        schema: TableSchema,
        frame: pd.DataFrame,
        geometry: dict[str, Any],
        pages: list[Any],
        crop_root: Path,
    ) -> list[ReviewCell]:
        panels = {panel["panel_id"]: panel for panel in geometry["panels"]}
        panel_for_variable = {
            column.variable: next(
                panel for panel in schema.panels if column.column_no in panel.printed_columns
            )
            for column in schema.get_all_columns()
        }
        child_variables = set(schema.hierarchy.child_variables) if schema.hierarchy else set()
        selections: list[tuple[str, int | None, int | None, int | None, pd.Series, str]] = []
        if schema.hierarchy is None:
            for _, row in frame.iterrows():
                for variable in schema.get_all_variables():
                    selections.append(
                        ("row", _integer(row["row_index"]), None, None, row, variable)
                    )
        else:
            for parent_index, group in frame.groupby("parent_row_index", sort=True):
                ordered_group = group.sort_values("subrow_index", key=lambda values: values.astype(int))
                parent_row = ordered_group.iloc[0]
                for variable in schema.get_all_variables():
                    if variable in child_variables:
                        for _, child_row in ordered_group.iterrows():
                            selections.append(
                                (
                                    "row",
                                    None,
                                    _integer(parent_index),
                                    _integer(child_row["subrow_index"]),
                                    child_row,
                                    variable,
                                )
                            )
                    else:
                        selections.append(
                            (
                                "parent",
                                None,
                                _integer(parent_index),
                                None,
                                parent_row,
                                variable,
                            )
                        )

        result: list[ReviewCell] = []
        for scope, row_index, parent_index, subrow_index, row, variable in selections:
            column = schema.get_column_by_var(variable)
            if column is None:
                raise ValueError(f"Unknown schema variable {variable!r}")
            panel_definition = panel_for_variable[variable]
            panel = panels[panel_definition.panel_id]
            x0, x1 = self._column_span(panel, column.column_no)
            y0, y1 = self._row_y_span(
                panel,
                schema,
                variable,
                row_index,
                parent_index,
                subrow_index,
            )
            page_number = int(panel["page"])
            selector_label = (
                f"r{row_index}"
                if row_index is not None
                else f"p{parent_index}" + (f"-s{subrow_index}" if subrow_index is not None else "")
            )
            crop_relative = (
                Path("crops")
                / document.pdf_id
                / selector_label
                / f"{column.column_no:02d}_{variable}.png"
            )
            crop_path = crop_root.parent / crop_relative
            crop_path.parent.mkdir(parents=True, exist_ok=True)
            page = pages[page_number - 1]
            clipped = (
                max(0, min(page.width, x0)),
                max(0, min(page.height, y0)),
                max(0, min(page.width, x1)),
                max(0, min(page.height, y1)),
            )
            if clipped[2] <= clipped[0] or clipped[3] <= clipped[1]:
                raise ValueError(f"{document.pdf_id}: empty crop for {selector_label}/{variable}")
            page.image.crop(clipped).save(crop_path, format="PNG")
            extracted = str(row[variable])
            review_key = "|".join(
                [
                    document.pdf_id,
                    scope,
                    "" if row_index is None else str(row_index),
                    "" if parent_index is None else str(parent_index),
                    "" if subrow_index is None else str(subrow_index),
                    variable,
                ]
            )
            review_id = hashlib.sha256(review_key.encode("utf-8")).hexdigest()[:20]
            result.append(
                ReviewCell(
                    review_id=review_id,
                    pdf_id=document.pdf_id,
                    district=document.district,
                    format_id=document.format_id,
                    scope=scope,
                    row_index=row_index,
                    parent_row_index=parent_index,
                    subrow_index=subrow_index,
                    variable=variable,
                    extracted_value=extracted,
                    review_status="PENDING",
                    corrected_value="",
                    reason="",
                    source_page=page_number,
                    panel_id=panel_definition.panel_id,
                    bbox=clipped,
                    crop_path=crop_relative.as_posix(),
                    crop_sha256=sha256_file(crop_path),
                )
            )
        return result

    @staticmethod
    def _expected_unique_cells(schema: TableSchema, frame: pd.DataFrame) -> int:
        if schema.hierarchy is None:
            return len(frame) * len(schema.get_all_variables())
        child_count = len(schema.hierarchy.child_variables)
        parent_count = frame["parent_row_index"].nunique()
        return parent_count * (len(schema.get_all_variables()) - child_count) + len(
            frame
        ) * child_count

    @staticmethod
    def _column_span(panel: dict[str, Any], column_no: int) -> tuple[int, int]:
        spans = panel.get("column_spans", {})
        entry = spans.get(str(column_no), spans.get(column_no))
        if entry:
            return int(entry["x_start"]), int(entry["x_end"])
        centers = {
            int(number): float(center) for number, center in panel["column_centers"].items()
        }
        ordered = sorted(centers)
        position = ordered.index(column_no)
        body_x0, _, body_x1, _ = panel["body_bbox"]
        x0 = (
            int(body_x0)
            if position == 0
            else round((centers[ordered[position - 1]] + centers[column_no]) / 2)
        )
        x1 = (
            int(body_x1)
            if position == len(ordered) - 1
            else round((centers[column_no] + centers[ordered[position + 1]]) / 2)
        )
        return x0, x1

    @staticmethod
    def _row_y_span(
        panel: dict[str, Any],
        schema: TableSchema,
        variable: str,
        row_index: int | None,
        parent_index: int | None,
        subrow_index: int | None,
    ) -> tuple[int, int]:
        if schema.hierarchy is not None and variable in schema.hierarchy.child_variables:
            hierarchy = panel.get("hierarchy")
            if hierarchy is None:
                raise ValueError("Hierarchy geometry is missing from the anchor panel")
            parent = next(
                item
                for item in hierarchy["parents"]
                if int(item["parent_row_index"]) == parent_index
            )
            subrow = next(
                item
                for item in parent["subrows"]
                if int(item["subrow_index"]) == subrow_index
            )
            return int(subrow["bbox"][1]), int(subrow["bbox"][3])
        target = parent_index if parent_index is not None else row_index
        row = next(item for item in panel["rows"] if int(item["row_index"]) == target)
        return int(row["bbox"][1]), int(row["bbox"][3])

    @staticmethod
    def _write_text_cell(cell: Any, value: Any) -> None:
        cell.value = "" if value is None else str(value)
        cell.data_type = "s"

    @classmethod
    def _write_workbook(
        cls,
        path: Path,
        cells: list[ReviewCell],
        table_summary: list[dict[str, Any]],
        *,
        manifest_sha: str,
        index_sha: str,
    ) -> None:
        workbook = Workbook()
        instructions = cast(Worksheet, workbook.active)
        instructions.title = "Instructions"
        summary = workbook.create_sheet("Summary")
        sheet = workbook.create_sheet("Cells")
        metadata = workbook.create_sheet("_Metadata")
        metadata.sheet_state = "hidden"

        navy = "17365D"
        teal = "0F6B78"
        pale_yellow = "FFF2CC"
        pale_green = "E2F0D9"
        pale_red = "FCE4D6"
        white = "FFFFFF"
        thin_gray = Side(style="thin", color="D9E2F3")

        instructions.sheet_view.showGridLines = False
        instructions["A1"] = "Remaining UP Source Review"
        instructions["A1"].font = Font(name="Aptos Display", size=18, bold=True, color=white)
        instructions["A1"].fill = PatternFill("solid", fgColor=navy)
        instructions.merge_cells("A1:F1")
        instructions["A3"] = "How to review"
        instructions["A3"].font = Font(size=12, bold=True, color=navy)
        steps = [
            "1. Open the Cells sheet and inspect each linked 300-DPI source crop.",
            "2. Set review_status to VERIFIED when extracted_value is exact.",
            "3. Set review_status to CORRECTED, enter corrected_value, and explain the change.",
            "4. Preserve printed spelling, punctuation, blanks, ellipses, dashes, and codes.",
            "5. Save this workbook and rerun main.py. Post-processing remains blocked until every row is reviewed.",
        ]
        for row_number, text in enumerate(steps, start=4):
            instructions.cell(row_number, 1, text)
        instructions["A11"] = "Do not edit the identifier, selector, source, bbox, crop, or hash columns."
        instructions["A11"].font = Font(bold=True, color="9C0006")
        instructions.column_dimensions["A"].width = 110

        summary.sheet_view.showGridLines = False
        summary["A1"] = "Source Review Summary"
        summary["A1"].font = Font(name="Aptos Display", size=16, bold=True, color=white)
        summary["A1"].fill = PatternFill("solid", fgColor=navy)
        summary.merge_cells("A1:F1")
        summary["A3"] = "Total unique cells"
        summary["A4"] = "Pending"
        summary["A5"] = "Verified"
        summary["A6"] = "Corrected"
        summary["B3"] = f"=COUNTA('Cells'!$A$2:$A${len(cells) + 1})"
        summary["B4"] = f'=COUNTIF(\'Cells\'!$K$2:$K${len(cells) + 1},"PENDING")'
        summary["B5"] = f'=COUNTIF(\'Cells\'!$K$2:$K${len(cells) + 1},"VERIFIED")'
        summary["B6"] = f'=COUNTIF(\'Cells\'!$K$2:$K${len(cells) + 1},"CORRECTED")'
        for row_number in range(3, 7):
            summary.cell(row_number, 1).font = Font(bold=True, color=navy)
            summary.cell(row_number, 1).fill = PatternFill("solid", fgColor="EAF2F8")
            summary.cell(row_number, 2).font = Font(bold=True)
        summary.append([])
        summary.append([])
        summary.append(["PDF", "District", "Format", "Rows", "Parents", "Unique cells"])
        for item in table_summary:
            summary.append(
                [
                    item["pdf_id"],
                    item["district"],
                    item["format_id"],
                    item["rows"],
                    item["parents"],
                    item["unique_cells"],
                ]
            )
        summary_header = summary[9]
        for cell in summary_header:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(bold=True, color=white)
        summary.freeze_panes = "A10"
        summary.auto_filter.ref = f"A9:F{summary.max_row}"
        for column, width in zip("ABCDEF", (34, 22, 14, 12, 12, 16), strict=True):
            summary.column_dimensions[column].width = width

        sheet.sheet_view.showGridLines = False
        sheet.append(REVIEW_HEADERS)
        for review in cells:
            sheet.append([None] * len(REVIEW_HEADERS))
            row_number = sheet.max_row
            for column_number, value in enumerate(review.workbook_row(), start=1):
                cls._write_text_cell(sheet.cell(row_number, column_number), value)
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=teal)
            cell.font = Font(bold=True, color=white)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin_gray)
        sheet.row_dimensions[1].height = 34
        sheet.freeze_panes = "A2"
        if cells:
            table = Table(displayName="SourceReviewCells", ref=f"A1:R{sheet.max_row}")
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
            sheet.add_table(table)
            status_column = REVIEW_HEADERS.index("review_status") + 1
            status_letter = sheet.cell(1, status_column).column_letter
            validation = DataValidation(
                type="list",
                formula1='"PENDING,VERIFIED,CORRECTED"',
                allow_blank=False,
            )
            validation.error = "Choose PENDING, VERIFIED, or CORRECTED."
            validation.errorTitle = "Invalid review status"
            sheet.add_data_validation(validation)
            validation.add(f"{status_letter}2:{status_letter}{sheet.max_row}")
            status_range = f"{status_letter}2:{status_letter}{sheet.max_row}"
            sheet.conditional_formatting.add(
                status_range,
                FormulaRule(
                    formula=[f'${status_letter}2="VERIFIED"'],
                    fill=PatternFill("solid", fgColor=pale_green),
                ),
            )
            sheet.conditional_formatting.add(
                status_range,
                FormulaRule(
                    formula=[f'${status_letter}2="CORRECTED"'],
                    fill=PatternFill("solid", fgColor=pale_yellow),
                ),
            )
            sheet.conditional_formatting.add(
                status_range,
                FormulaRule(
                    formula=[f'${status_letter}2="PENDING"'],
                    fill=PatternFill("solid", fgColor=pale_red),
                ),
            )
            crop_column = REVIEW_HEADERS.index("crop_path") + 1
            for row_number in range(2, sheet.max_row + 1):
                crop_cell = sheet.cell(row_number, crop_column)
                crop_cell.hyperlink = crop_cell.value
                crop_cell.style = "Hyperlink"

        widths = {
            "A": 22,
            "B": 32,
            "C": 20,
            "D": 14,
            "E": 10,
            "F": 11,
            "G": 16,
            "H": 14,
            "I": 32,
            "J": 34,
            "K": 16,
            "L": 34,
            "M": 38,
            "N": 12,
            "O": 26,
            "P": 22,
            "Q": 70,
            "R": 66,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)

        metadata.append(["key", "value"])
        metadata.append(["source_manifest_sha256", manifest_sha])
        metadata.append(["review_index_sha256", index_sha])
        metadata.append(["unique_cells", len(cells)])
        try:
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
        except AttributeError:
            pass
        temporary = path.with_name(f".{path.name}.{time.time_ns()}.tmp")
        workbook.save(temporary)
        temporary.replace(path)


class ReviewWorkbookCompiler:
    """Validate every review row and compile an explicit version-2 ledger."""

    def compile(
        self,
        *,
        workbook_path: Path,
        index_path: Path,
        ledger_path: Path,
    ) -> ReviewCompilationResult | None:
        workbook_path = Path(workbook_path).resolve()
        index_path = Path(index_path).resolve()
        ledger_path = Path(ledger_path).resolve()
        index = json.loads(index_path.read_text(encoding="utf-8"))
        manifest_path = Path(index["source_manifest"])
        if sha256_file(manifest_path) != index["source_manifest_sha256"]:
            raise ValueError("Extraction manifest changed after review package generation")
        workbook = load_workbook(workbook_path, data_only=False, read_only=False)
        try:
            if "Cells" not in workbook.sheetnames or "_Metadata" not in workbook.sheetnames:
                raise ValueError("Review workbook is missing required sheets")
            metadata = {
                str(row[0].value): str(row[1].value)
                for row in workbook["_Metadata"].iter_rows(min_row=2, max_col=2)
                if row[0].value is not None
            }
            if metadata.get("source_manifest_sha256") != index["source_manifest_sha256"]:
                raise ValueError("Workbook source-manifest hash was changed")
            if metadata.get("review_index_sha256") != sha256_file(index_path):
                raise ValueError("Workbook review-index hash was changed")
            sheet = workbook["Cells"]
            headers = [str(cell.value or "") for cell in sheet[1]]
            if headers != REVIEW_HEADERS:
                raise ValueError("Review workbook headers were changed")
            expected = index["cells"]
            actual_rows = list(sheet.iter_rows(min_row=2, values_only=False))
            if len(actual_rows) != len(expected):
                raise ValueError(
                    f"Review row count {len(actual_rows)} != expected {len(expected)}"
                )

            ledger_tables: dict[str, dict[str, Any]] = {}
            corrected_count = 0
            pending = 0
            seen: set[str] = set()
            for excel_row, expected_row in zip(actual_rows, expected, strict=True):
                actual = {
                    header: self._value(excel_row[index].value)
                    for index, header in enumerate(REVIEW_HEADERS)
                }
                expected_values = dict(expected_row)
                expected_values["bbox"] = ",".join(
                    str(value) for value in expected_values["bbox"]
                )
                for header in IMMUTABLE_HEADERS:
                    if actual[header] != self._value(expected_values[header]):
                        raise ValueError(
                            f"{expected_row['review_id']}: immutable field {header!r} changed"
                        )
                review_id = actual["review_id"]
                if review_id in seen:
                    raise ValueError(f"Duplicate review id {review_id}")
                seen.add(review_id)
                crop_path = index_path.parent / expected_row["crop_path"]
                if not crop_path.is_file() or sha256_file(crop_path) != expected_row["crop_sha256"]:
                    raise ValueError(f"{review_id}: review crop is missing or changed")
                status = actual["review_status"].strip().upper()
                if status == "PENDING" or not status:
                    pending += 1
                    continue
                if status not in {"VERIFIED", "CORRECTED"}:
                    raise ValueError(f"{review_id}: invalid review status {status!r}")
                original = expected_row["extracted_value"]
                corrected = actual["corrected_value"]
                reason = actual["reason"].strip()
                if status == "VERIFIED":
                    if corrected:
                        raise ValueError(f"{review_id}: VERIFIED row cannot contain a correction")
                    verified_value = original
                else:
                    verified_value = corrected
                    if verified_value == original:
                        raise ValueError(f"{review_id}: CORRECTED value did not change")
                    if not reason:
                        raise ValueError(f"{review_id}: CORRECTED row requires a reason")
                    corrected_count += 1
                pdf_id = expected_row["pdf_id"]
                table = ledger_tables.setdefault(
                    pdf_id,
                    {
                        "format_id": expected_row["format_id"],
                        "expected_rows": 0,
                        "expected_parent_rows": 0,
                        "reviewed_unique_cells": 0,
                        "reviews": [],
                    },
                )
                selector = self._selector(expected_row)
                table["reviews"].append(
                    {
                        "selector": selector,
                        "scope": expected_row["scope"],
                        "variable": expected_row["variable"],
                        "expected_original": original,
                        "verified_value": verified_value,
                        "status": status,
                        "source_page": expected_row["source_page"],
                        "panel_id": expected_row["panel_id"],
                        "bbox": expected_row["bbox"],
                        "crop_sha256": expected_row["crop_sha256"],
                        "reason": reason,
                    }
                )
            if pending:
                return None

            summary = {item["pdf_id"]: item for item in index["table_summary"]}
            for pdf_id, table in ledger_tables.items():
                item = summary[pdf_id]
                table["expected_rows"] = item["rows"]
                table["expected_parent_rows"] = item["parents"]
                table["reviewed_unique_cells"] = item["unique_cells"]
            ledger = {
                "version": 2,
                "source_run_id": index["source_run_id"],
                "source_manifest_sha256": index["source_manifest_sha256"],
                "expected_table_count": len(ledger_tables),
                "expected_output_rows": sum(item["rows"] for item in summary.values()),
                "expected_unique_source_cells": len(expected),
                "tables": ledger_tables,
            }
            CorrectionLedger.model_validate(ledger)
            ledger_path.parent.mkdir(parents=True, exist_ok=True)
            _atomic_text(
                ledger_path,
                yaml.safe_dump(ledger, sort_keys=False, allow_unicode=True),
            )
            return ReviewCompilationResult(ledger_path, len(expected), corrected_count)
        finally:
            workbook.close()

    @staticmethod
    def _value(value: Any) -> str:
        if value is None:
            return ""
        return str(value)

    @staticmethod
    def _selector(row: dict[str, Any]) -> dict[str, int]:
        if row["row_index"] is not None:
            return {"row_index": int(row["row_index"])}
        selector = {"parent_row_index": int(row["parent_row_index"])}
        if row["subrow_index"] is not None:
            selector["subrow_index"] = int(row["subrow_index"])
        return selector
